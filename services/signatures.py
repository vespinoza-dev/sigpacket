"""
Signature log (Excel) management Blueprint.
Extracted from app.py — handles saving, reading, deleting, and downloading
signature entries stored in an Excel workbook.
"""

from flask import Blueprint, jsonify, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.exceptions import InvalidFileException
import json
import os
import tempfile
from zipfile import BadZipFile

signatures_bp = Blueprint('signatures', __name__)

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

SIGNATURES_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "signatures.xlsx")
SIGNATURES_DIR = os.path.dirname(SIGNATURES_FILE)


def _safe_load():
    """Load workbook, auto-recovering from corruption."""
    try:
        return load_workbook(SIGNATURES_FILE)
    except (BadZipFile, InvalidFileException, KeyError):
        os.remove(SIGNATURES_FILE)
        ensure_excel()
        return load_workbook(SIGNATURES_FILE)


def _safe_save(wb):
    """Write workbook to a temp file then atomically replace the target."""
    fd, tmp_path = tempfile.mkstemp(suffix=".xlsx", dir=SIGNATURES_DIR)
    try:
        os.close(fd)
        wb.save(tmp_path)
        os.replace(tmp_path, SIGNATURES_FILE)
    except Exception:
        # Clean up temp file on failure
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
        raise

EXCEL_COLUMNS = ["Client", "Signer Type", "Signer Name", "Title", "Signing Entity",
                 "Additional Signing Entity", "Additional Entity Role",
                 "Additional Signing Entity 2", "Additional Entity Role 2",
                 "Additional Signing Entity 3", "Additional Entity Role 3",
                 "Email", "Phone", "CC Email",
                 "Address", "City State ZIP", "Signer Description", "Field Overrides"]

# ---------------------------------------------------------------------------
# Data functions
# ---------------------------------------------------------------------------


def ensure_excel():
    """Create the Excel file with headers if it doesn't exist."""
    if os.path.exists(SIGNATURES_FILE):
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Signatures"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1A2744", end_color="1A2744", fill_type="solid")
    thin_border = Border(
        bottom=Side(style="thin", color="D1D5DB")
    )

    for col_idx, col_name in enumerate(EXCEL_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    # Set column widths
    widths = [20, 14, 22, 28, 28, 28, 28, 28, 28, 28, 28, 28, 18, 28, 28, 28, 40, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    _safe_save(wb)


def add_signature_row(client, fields):
    """Append a signature entry to the Excel log."""
    ensure_excel()
    wb = _safe_load()
    ws = wb.active
    row = _build_signature_row(client, fields)
    ws.append(row)
    _safe_save(wb)


def get_all_signatures():
    """Read all signature entries from the Excel log."""
    ensure_excel()
    try:
        wb = _safe_load()
    except (BadZipFile, InvalidFileException, KeyError):
        # File is corrupted — recreate it
        os.remove(SIGNATURES_FILE)
        ensure_excel()
        return []
    ws = wb.active
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(cell is not None for cell in row):
            rows.append({
                "client": row[0] or "",
                "signer_type": (row[1] or "entity") if len(row) > 1 else "entity",
                "signer_name": (row[2] or "") if len(row) > 2 else "",
                "title": (row[3] or "") if len(row) > 3 else "",
                "signing_entity": (row[4] or "") if len(row) > 4 else "",
                "additional_signing_entity": (row[5] or "") if len(row) > 5 else "",
                "additional_signing_entity_title": (row[6] or "") if len(row) > 6 else "",
                "additional_signing_entity_2": (row[7] or "") if len(row) > 7 else "",
                "additional_signing_entity_title_2": (row[8] or "") if len(row) > 8 else "",
                "additional_signing_entity_3": (row[9] or "") if len(row) > 9 else "",
                "additional_signing_entity_title_3": (row[10] or "") if len(row) > 10 else "",
                "email": (row[11] or "") if len(row) > 11 else "",
                "phone": (row[12] or "") if len(row) > 12 else "",
                "cc_email": (row[13] or "") if len(row) > 13 else "",
                "address": (row[14] or "") if len(row) > 14 else "",
                "city_state_zip": (row[15] or "") if len(row) > 15 else "",
                "signer_description": (row[16] or "") if len(row) > 16 else "",
                "field_overrides": json.loads(row[17]) if len(row) > 17 and row[17] else {},
            })
    return rows


def update_signature_row(index, client, fields):
    """Update a signature row by index (0-based, excluding header)."""
    ensure_excel()
    wb = _safe_load()
    ws = wb.active
    row_num = index + 2  # 1-based + skip header
    if row_num > ws.max_row:
        return False
    if client is not None:
        ws.cell(row=row_num, column=1, value=client)
    col_map = {
        "signer_type": 2,
        "signer_name": 3,
        "title": 4,
        "signing_entity": 5,
        "additional_signing_entity": 6,
        "additional_signing_entity_title": 7,
        "additional_signing_entity_2": 8,
        "additional_signing_entity_title_2": 9,
        "additional_signing_entity_3": 10,
        "additional_signing_entity_title_3": 11,
        "email": 12,
        "phone": 13,
        "cc_email": 14,
        "address": 15,
        "city_state_zip": 16,
        "signer_description": 17,
    }
    for key, col in col_map.items():
        if key in fields:
            ws.cell(row=row_num, column=col, value=fields[key])
    # Handle field_overrides separately (JSON serialization)
    if "field_overrides" in fields:
        ws.cell(row=row_num, column=18, value=json.dumps(fields["field_overrides"]) if fields["field_overrides"] else "")
    _safe_save(wb)
    return True


def delete_signature_row(index):
    """Delete a signature row by index (0-based, excluding header)."""
    ensure_excel()
    wb = _safe_load()
    ws = wb.active
    row_num = index + 2  # 1-based + skip header
    if row_num <= ws.max_row:
        ws.delete_rows(row_num)
        _safe_save(wb)
        return True
    return False


def clear_all_signatures():
    """Delete the workbook and recreate it fresh."""
    if os.path.exists(SIGNATURES_FILE):
        os.remove(SIGNATURES_FILE)
    ensure_excel()


def _build_signature_row(client, fields):
    """Convert one signature record into workbook row values."""
    return [
        client,
        fields.get("signer_type", "entity"),
        fields.get("signer_name", ""),
        fields.get("title", ""),
        fields.get("signing_entity", ""),
        fields.get("additional_signing_entity", ""),
        fields.get("additional_signing_entity_title", ""),
        fields.get("additional_signing_entity_2", ""),
        fields.get("additional_signing_entity_title_2", ""),
        fields.get("additional_signing_entity_3", ""),
        fields.get("additional_signing_entity_title_3", ""),
        fields.get("email", ""),
        fields.get("phone", ""),
        fields.get("cc_email", ""),
        fields.get("address", ""),
        fields.get("city_state_zip", ""),
        fields.get("signer_description", ""),
        json.dumps(fields.get("field_overrides", {})) if fields.get("field_overrides") else "",
    ]


def _snapshot_signature_entries():
    """Capture the current workbook rows in re-importable form."""
    return [
        {
            "client": row.get("client", ""),
            "fields": {
                "signer_type": row.get("signer_type", "entity"),
                "signer_name": row.get("signer_name", ""),
                "title": row.get("title", ""),
                "signing_entity": row.get("signing_entity", ""),
                "additional_signing_entity": row.get("additional_signing_entity", ""),
                "additional_signing_entity_title": row.get("additional_signing_entity_title", ""),
                "additional_signing_entity_2": row.get("additional_signing_entity_2", ""),
                "additional_signing_entity_title_2": row.get("additional_signing_entity_title_2", ""),
                "additional_signing_entity_3": row.get("additional_signing_entity_3", ""),
                "additional_signing_entity_title_3": row.get("additional_signing_entity_title_3", ""),
                "email": row.get("email", ""),
                "phone": row.get("phone", ""),
                "cc_email": row.get("cc_email", ""),
                "address": row.get("address", ""),
                "city_state_zip": row.get("city_state_zip", ""),
                "signer_description": row.get("signer_description", ""),
                "field_overrides": row.get("field_overrides", {}),
            },
        }
        for row in get_all_signatures()
    ]


def _replace_all_signatures(entries):
    """Replace workbook contents with the provided signature entries."""
    ensure_excel()
    wb = _safe_load()
    ws = wb.active

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for entry in entries:
        ws.append(_build_signature_row(entry.get("client", ""), entry.get("fields", {})))

    _safe_save(wb)


# ---------------------------------------------------------------------------
# Route handlers
# ---------------------------------------------------------------------------

@signatures_bp.route("/signatures", methods=["GET"])
def get_signatures():
    """Return all saved signatures as JSON."""
    return jsonify({"signatures": get_all_signatures()})


@signatures_bp.route("/signatures", methods=["POST"])
def save_signature():
    """Save extracted fields to the signatures Excel log."""
    payload = request.get_json()
    client = payload.get("client", "").strip()
    fields = payload.get("fields", {})
    if not fields.get("signer_name", "").strip() and not fields.get("signing_entity", "").strip():
        return jsonify({"error": "At least a signer name or signing entity is required"}), 400
    add_signature_row(client, fields)
    return jsonify({"ok": True})


@signatures_bp.route("/signatures/<int:index>", methods=["PUT"])
def update_signature(index):
    """Update a signature row by index."""
    payload = request.get_json()
    fields = payload.get("fields", {})
    client = payload.get("client")
    if update_signature_row(index, client, fields):
        return jsonify({"ok": True})
    return jsonify({"error": "Row not found"}), 404


@signatures_bp.route("/signatures/<int:index>", methods=["DELETE"])
def delete_signature(index):
    """Delete a signature row by index."""
    if delete_signature_row(index):
        return jsonify({"ok": True})
    return jsonify({"error": "Row not found"}), 404


@signatures_bp.route("/signatures/upload", methods=["POST"])
def upload_signatures():
    """Upload an Excel file to populate signatories and document assignments."""
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    file = request.files["file"]
    if not file.filename.endswith(".xlsx"):
        return jsonify({"error": "Only .xlsx files are supported"}), 400

    try:
        wb = load_workbook(file)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        return jsonify({"error": f"Invalid workbook: {exc}"}), 400

    col_map = {
        "Client": "client",
        "Signer Type": "signer_type",
        "Signer Name": "signer_name",
        "Title": "title",
        "Signing Entity": "signing_entity",
        "Additional Signing Entity": "additional_signing_entity",
        "Additional Entity Role": "additional_signing_entity_title",
        "Additional Signing Entity 2": "additional_signing_entity_2",
        "Additional Entity Role 2": "additional_signing_entity_title_2",
        "Additional Signing Entity 3": "additional_signing_entity_3",
        "Additional Entity Role 3": "additional_signing_entity_title_3",
        "Email": "email",
        "Phone": "phone",
        "CC Email": "cc_email",
        "Address": "address",
        "City State ZIP": "city_state_zip",
        "Signer Description": "signer_description",
    }

    try:
        # --- Parse Signatories sheet ---
        sig_sheet = None
        for name in ["Signatories", "Sheet1"]:
            if name in wb.sheetnames:
                sig_sheet = wb[name]
                break
        if sig_sheet is None:
            if not wb.worksheets:
                return jsonify({"error": "Workbook has no worksheets"}), 400
            sig_sheet = wb.worksheets[0]

        header_rows = sig_sheet.iter_rows(min_row=1, max_row=1, values_only=True)
        headers = list(next(header_rows, ()))
        if not headers:
            return jsonify({"error": "Signatories sheet is missing a header row"}), 400

        imported_entries = []
        signatories = []
        for row in sig_sheet.iter_rows(min_row=2, values_only=True):
            if not any(cell is not None for cell in row):
                continue
            fields = {}
            client = ""
            for i, header in enumerate(headers):
                if i >= len(row):
                    break
                if header in col_map:
                    key = col_map[header]
                    val = str(row[i]) if row[i] is not None else ""
                    if key == "client":
                        client = val
                    else:
                        fields[key] = val
            if fields.get("signer_name") or fields.get("signing_entity"):
                imported_entries.append({"client": client, "fields": fields})
                signatories.append(fields)

        if not imported_entries:
            return jsonify({"error": "No usable signatories found in the workbook"}), 400

        # --- Parse Document Assignments sheet ---
        doc_assignments = {}
        if "Document Assignments" in wb.sheetnames:
            assign_sheet = wb["Document Assignments"]
            assign_headers = list(next(assign_sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
            doc_type_labels = assign_headers[1:]

            for row_idx, row in enumerate(assign_sheet.iter_rows(min_row=2, values_only=True)):
                if not any(cell is not None for cell in row):
                    continue
                checked = []
                for col_idx, label in enumerate(doc_type_labels):
                    val = row[col_idx + 1] if col_idx + 1 < len(row) else None
                    if val and str(val).strip().upper() in ("X", "TRUE", "YES"):
                        checked.append(label)
                doc_assignments[str(row_idx)] = checked
    except Exception as exc:
        return jsonify({"error": f"Failed to parse uploaded workbook: {exc}"}), 400

    previous_entries = _snapshot_signature_entries()
    try:
        _replace_all_signatures(imported_entries)
    except Exception:
        try:
            _replace_all_signatures(previous_entries)
        except Exception:
            pass
        return jsonify({"error": "Failed to save uploaded signatures"}), 500

    return jsonify({"ok": True, "count": len(signatories), "doc_assignments": doc_assignments})


@signatures_bp.route("/signatures/download", methods=["GET", "POST"])
def download_signatures():
    """Download the signatures Excel file with per-document columns."""
    ensure_excel()

    if request.method == "POST":
        payload = request.get_json() or {}
        matrix_selections = payload.get("matrix_selections", {})
        doc_columns = payload.get("doc_columns", [])

        source_wb = _safe_load()
        source_ws = source_wb.active

        wb = Workbook()
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1A2744", end_color="1A2744", fill_type="solid")
        thin_border = Border(bottom=Side(style="thin", color="D1D5DB"))

        # --- Sheet 1: Signatories ---
        ws1 = wb.active
        ws1.title = "Signatories"

        sig_columns = [c for c in EXCEL_COLUMNS if c != "Field Overrides"]
        for col_idx, col_name in enumerate(sig_columns, 1):
            cell = ws1.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        num_sig_cols = len(sig_columns)
        signer_names = []
        for row in source_ws.iter_rows(min_row=2, values_only=True):
            if not any(cell is not None for cell in row):
                continue
            row_data = list(row[:num_sig_cols])
            while len(row_data) < num_sig_cols:
                row_data.append("")
            ws1.append(row_data)
            # Track identifier for Sheet 2: entity name for entities, signer name for individuals
            signer_type = (row_data[1] or "entity").lower()
            if signer_type == "individual":
                signer_names.append(row_data[2] or "")  # Signer Name
            else:
                signer_names.append(row_data[4] or row_data[2] or "")  # Signing Entity, fallback to Signer Name

        sig_widths = [20, 14, 22, 28, 28, 28, 28, 28, 28, 28, 28, 28, 18, 28, 28, 28, 40]
        for i, w in enumerate(sig_widths, 1):
            if i <= num_sig_cols:
                ws1.column_dimensions[ws1.cell(row=1, column=i).column_letter].width = w

        # --- Sheet 2: Document Assignments ---
        ws2 = wb.create_sheet("Document Assignments")

        doc_labels = [d["label"] for d in doc_columns]
        doc_keys = [d["key"] for d in doc_columns]
        assign_columns = ["Name"] + doc_labels

        for col_idx, col_name in enumerate(assign_columns, 1):
            cell = ws2.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        for row_idx, name in enumerate(signer_names):
            selected = set(matrix_selections.get(str(row_idx), []))
            row_data = [name]
            for key in doc_keys:
                row_data.append("X" if key in selected else "")
            ws2.append(row_data)

        assign_widths = [28] + [22] * len(doc_labels)
        for i, w in enumerate(assign_widths, 1):
            if i <= len(assign_columns):
                ws2.column_dimensions[ws2.cell(row=1, column=i).column_letter].width = w

        from io import BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True,
                         download_name="signatures.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    # Plain GET — serve raw file as before
    return send_file(SIGNATURES_FILE, as_attachment=True,
                     download_name="signatures.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
